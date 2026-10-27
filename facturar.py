"""
Panel «Por facturar» del módulo de Ventas.

Muestra todos los pedidos que se guardaron con «Este pedido lleva DATOS DE FACTURACIÓN» (pendientes y ya
facturados; estos últimos solo de consulta) y arma las partidas tal como se capturan en CONTPAQi (cantidad, descripción, unidad,
precio unitario sin IVA, importe e impuesto), con cada dato listo para copiar.

Se pueden elegir varios pedidos a la vez (factura mostrador). Al final se captura el N° de factura una sola vez
y queda guardado en todos los pedidos elegidos (con eso salen de la lista).

No necesita tablas nuevas: usa pedidos.requiere_facturacion, concepto, uso_cfdi, no_factura y pedido_items.
"""
import datetime as dt
import html
from decimal import Decimal, ROUND_HALF_UP

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

import codigos

# Lo que se muestra en las columnas UNIDAD e IMPUESTO (cámbialo si en CONTPAQi usas otros nombres)
UNIDAD = "Pieza"
IMPUESTO = "IVA 16%"
TASA_IVA = Decimal("0.16")
DESGLOSAR = "Desglosar productos"
DESCRIPCION_ENVIO = "SERVICIO DE ENVÍO"
CONCEPTOS = ["Accesorios de decoración", "Mobiliario de oficina", DESGLOSAR]
from permisos import roles as _roles
ROLES_FACTURAR = _roles("por_facturar")


# ----------------------------------------------------------------------
# Cálculos (sin Streamlit, para poder probarlos)
# ----------------------------------------------------------------------
def _d(x):
    return Decimal(str(x or 0))


def _r2(x):
    return Decimal(x).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def importe_sin_iva(total_con_iva):
    """Importe sin IVA tal que importe + IVA(16%, redondeado a centavos) dé EXACTAMENTE el total cobrado.
    Devuelve (importe, iva, cuadra)."""
    total = _r2(_d(total_con_iva))
    base = _r2(total / (1 + TASA_IVA))
    for ajuste in ("0", "-0.01", "0.01", "-0.02", "0.02"):
        imp = base + Decimal(ajuste)
        iva = _r2(imp * TASA_IVA)
        if imp + iva == total:
            return imp, iva, True
    # ~1 de cada 7 totales no tiene un precio de 2 decimales que cuadre: se da el precio con 6 decimales
    imp6 = (total / (1 + TASA_IVA)).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP)
    return imp6, total - imp6, False


def precio_unitario(importe, cantidad):
    """Precio unitario sin IVA. Con cantidad 1 es el importe; si no divide exacto se dan hasta 6 decimales."""
    cantidad = int(cantidad or 1)
    if cantidad <= 1:
        return importe, True
    exacto = _r2(importe / cantidad)
    if exacto * cantidad == importe:
        return exacto, True
    return (importe / cantidad).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP), False


def armar_partidas(pedidos, items_por_pedido, concepto_de, con_modelo=False):
    """Agrupa los renglones de los pedidos elegidos según el concepto de cada pedido:
    - «Accesorios de decoración» / «Mobiliario de oficina»: una sola partida por concepto con la suma.
    - «Desglosar productos»: una partida por pieza.
    Devuelve lista de dicts con cantidad, descripcion, unidad, precio_unitario, importe, iva, total, impuesto, avisos."""
    grupos, desglose = {}, []
    for ped in pedidos:
        num = ped["numero_pedido"]
        concepto = concepto_de.get(num) or DESGLOSAR
        for it in items_por_pedido.get(num, []):
            total_linea = _d(it.get("total_linea"))
            if concepto == DESGLOSAR:
                desc = (it.get("descripcion") or "").strip() or "Sin descripción"
                if con_modelo and it.get("modelo"):
                    desc = f"{it['modelo']} - {desc}"
                desglose.append({"cantidad": int(it.get("cantidad") or 1), "descripcion": desc, "total": total_linea})
            else:
                grupos[concepto] = grupos.get(concepto, Decimal("0")) + total_linea
    filas = [{"cantidad": 1, "descripcion": c.upper(), "total": t} for c, t in grupos.items()] + desglose
    envio = sum((_d(p.get("envio") or 0) for p in pedidos), Decimal("0"))
    if envio > 0:      # el envío va como su propia partida (si son varios pedidos, una sola con la suma)
        filas.append({"cantidad": 1, "descripcion": DESCRIPCION_ENVIO, "total": envio})

    partidas = []
    for f in filas:
        if f["total"] <= 0:
            continue
        imp, iva, cuadra = importe_sin_iva(f["total"])
        pu, pu_exacto = precio_unitario(imp, f["cantidad"])
        avisos = []
        if not cuadra:
            avisos.append("Precio con 6 decimales para que el total cuadre al centavo (cópialo completo).")
        if not pu_exacto and cuadra:
            avisos.append("El precio unitario no sale exacto: mejor captura cantidad 1 con el importe y pon la "
                          "cantidad en la descripción.")
        partidas.append({
            "cantidad": f["cantidad"], "descripcion": f["descripcion"], "unidad": UNIDAD,
            "precio_unitario": pu, "importe": imp, "iva": iva, "total": _r2(f["total"]),
            "impuesto": IMPUESTO, "avisos": avisos,
        })
    return partidas


def _fmt(x, dec=2):
    """Número sin $ ni comas, como se escribe en CONTPAQi."""
    s = f"{Decimal(x):.{dec}f}"
    if dec > 2:
        s = s.rstrip("0")
        if len(s.split(".")[1]) < 2:
            s = f"{Decimal(x):.2f}"
    return s


def _fmt_unitario(x):
    return _fmt(x, 6) if Decimal(x) != _r2(x) else _fmt(x)


# ----------------------------------------------------------------------
# Tabla HTML con «clic para copiar»
# ----------------------------------------------------------------------
_CSS = """
<style>
 body{margin:0;font-family:"Source Sans Pro",-apple-system,Segoe UI,Roboto,sans-serif;color:#31333F;background:transparent}
 table{border-collapse:collapse;width:100%;font-size:14px}
 th{text-align:left;font-weight:600;font-size:12px;color:#6b6d76;padding:6px 8px;border-bottom:1px solid #ddd;white-space:nowrap}
 td{padding:6px 8px;border-bottom:1px solid #eee;vertical-align:top}
 td.c{cursor:pointer;border-radius:4px}
 td.c:hover{background:#eef3ff}
 td.ok{background:#d9f5df !important}
 td.n{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
 tr.tot td{font-weight:600;border-top:2px solid #ccc}
 .bar{display:flex;gap:8px;align-items:center;margin:0 0 8px 0;font-size:13px;color:#6b6d76}
 button{font:inherit;font-size:13px;padding:4px 10px;border:1px solid #ccc;border-radius:6px;background:#fff;cursor:pointer}
 button:hover{background:#f3f3f3}
 .av{color:#a35a00;font-size:12px}
 @media (prefers-color-scheme: dark){body{color:#fafafa}th{color:#aaa;border-color:#444}td{border-color:#333}
   td.c:hover{background:#26324d}td.ok{background:#1f4d2b !important}button{background:#262730;color:#fafafa;border-color:#555}}
</style>
"""

_JS = """
<script>
function copiar(t){
  if(navigator.clipboard&&window.isSecureContext){return navigator.clipboard.writeText(t).catch(()=>viejo(t));}
  viejo(t);return Promise.resolve();
}
function viejo(t){const a=document.createElement('textarea');a.value=t;document.body.appendChild(a);a.select();
  try{document.execCommand('copy')}catch(e){}document.body.removeChild(a);}
document.querySelectorAll('td.c').forEach(td=>td.addEventListener('click',()=>{
  copiar(td.dataset.v);document.querySelectorAll('td.ok').forEach(x=>x.classList.remove('ok'));td.classList.add('ok');
}));
document.querySelectorAll('button[data-v]').forEach(b=>b.addEventListener('click',()=>{
  copiar(b.dataset.v);const o=b.textContent;b.textContent='¡Copiado!';setTimeout(()=>b.textContent=o,1200);
}));
</script>
"""


def _celda(valor, mostrar=None, num=False, extra=""):
    """`extra` es HTML ya escapado que se agrega debajo del dato (avisos)."""
    v = html.escape(str(valor), quote=True)
    m = html.escape(str(mostrar if mostrar is not None else valor)) + extra
    return f'<td class="c{" n" if num else ""}" data-v="{v}" title="Clic para copiar">{m}</td>'


def html_partidas(partidas):
    filas, tsv = [], ["Cantidad\tDescripción\tUnidad\tPrecio unitario\tImporte\tImpuesto"]
    for p in partidas:
        pu = _fmt_unitario(p["precio_unitario"])
        av = "".join(f'<div class="av">⚠ {html.escape(a)}</div>' for a in p["avisos"])
        filas.append(
            "<tr>" + _celda(p["cantidad"], num=True)
            + _celda(p["descripcion"], extra=av)
            + _celda(p["unidad"]) + _celda(pu, f"{Decimal(pu):,.{max(2, len(pu.split('.')[1]))}f}", num=True)
            + _celda(_fmt(p["importe"]), f"{p['importe']:,.2f}", num=True) + _celda(p["impuesto"])
            + f'<td class="n">{p["iva"]:,.2f}</td><td class="n">{p["total"]:,.2f}</td></tr>')
        tsv.append(f"{p['cantidad']}\t{p['descripcion']}\t{p['unidad']}\t{pu}\t{_fmt(p['importe'])}\t{p['impuesto']}")
    sub = sum((p["importe"] for p in partidas), Decimal("0"))
    iva = sum((p["iva"] for p in partidas), Decimal("0"))
    tot = sum((p["total"] for p in partidas), Decimal("0"))
    filas.append(f'<tr class="tot"><td></td><td>TOTALES</td><td></td><td></td>'
                 f'<td class="n">{sub:,.2f}</td><td></td><td class="n">{iva:,.2f}</td><td class="n">{tot:,.2f}</td></tr>')
    cuerpo = (
        f'<div class="bar"><span>Clic en cualquier dato para copiarlo.</span>'
        f'<button data-v="{html.escape(chr(10).join(tsv), quote=True)}">Copiar tabla completa</button></div>'
        "<table><tr><th>CANTIDAD</th><th>DESCRIPCIÓN</th><th>UNIDAD</th><th>PRECIO UNIT. (sin IVA)</th>"
        "<th>IMPORTE</th><th>IMPUESTO</th><th>IVA</th><th>TOTAL</th></tr>" + "".join(filas) + "</table>")
    return _CSS + cuerpo + _JS


def excel_contpaqi(partidas):
    """Excel para «Importar/Exportar Excel» de la factura en CONTPAQi (Cantidad, Descripción, Unidad, Precio Unitario).
    CONTPAQi los importa en orden inverso, así que se escriben al revés para que queden en el orden de la app.
    El IVA 16% se elige en CONTPAQi en cada renglón después de importar."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Partidas"
    ws.append(["Cantidad", "Descripción", "Unidad", "Precio Unitario"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for p in reversed(partidas):
        ws.append([int(p["cantidad"]), p["descripcion"], p["unidad"], float(p["precio_unitario"])])
    for fila in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in fila:
            c.number_format = "0.00####"
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["D"].width = 16
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def html_datos(pares):
    filas = "".join(f'<tr><th style="width:170px">{html.escape(k)}</th>{_celda(v)}</tr>' for k, v in pares if v)
    if not filas:
        return ""
    return _CSS + '<div class="bar"><span>Clic en un dato para copiarlo.</span></div><table>' + filas + "</table>" + _JS


# ----------------------------------------------------------------------
# Base de datos
# ----------------------------------------------------------------------
def facturado(p):
    return bool((p.get("no_factura") or "").strip())


from caches import guardar_en_memoria


def pedidos_con_factura(sb, desde):
    return _pedidos_con_factura_mem(sb, desde)


@guardar_en_memoria(ttl=30)
def _pedidos_con_factura_mem(_sb, desde):
    return _pedidos_con_factura(_sb, desde)


def _pedidos_con_factura(sb, desde):
    """Todos los pedidos que llevan facturación (facturados o no), sin los cancelados."""
    columnas = ("numero_pedido,fecha,cliente,total,concepto,uso_cfdi,email,telefono,email_facturacion,"
                "telefono_facturacion,notas_facturacion,status,estado_pedido,efectivo,no_factura")
    def leer(cols):
        return (sb.table("pedidos").select(cols)
                .eq("requiere_facturacion", True).gte("fecha", desde.isoformat())
                .order("numero_pedido", desc=True).limit(1000).execute().data)
    try:
        rows = leer(columnas + ",envio")
    except Exception:          # todavía no existe la columna «envio»
        rows = leer(columnas)
    return [r for r in rows if r.get("status") != "Cancelado" and r.get("estado_pedido") != "Cancelado"]


def items_de(sb, numeros):
    if not numeros:
        return {}
    return _items_de_mem(sb, tuple(sorted(numeros)))


@guardar_en_memoria(ttl=30)
def _items_de_mem(_sb, numeros):
    rows = (_sb.table("pedido_items").select("numero_pedido,modelo,descripcion,cantidad,total_linea")
            .in_("numero_pedido", list(numeros)).order("id").execute().data)
    out = {}
    for r in rows:
        out.setdefault(r["numero_pedido"], []).append(r)
    return out


def marcar_facturados(sb, pedidos, no_factura):
    """Guarda el N° de factura en cada pedido y su total como «facturado con IVA» (sin tocar el efectivo).
    Solo escribe en pedidos que siguen sin N° de factura (si alguien más ya lo facturó, no se sobrescribe).
    Devuelve (guardados, ya_tenian)."""
    guardados, ya_tenian = [], []
    for p in pedidos:
        num = p["numero_pedido"]
        actual = sb.table("pedidos").select("no_factura").eq("numero_pedido", num).execute().data
        if actual and (actual[0].get("no_factura") or "").strip():
            ya_tenian.append((num, actual[0]["no_factura"].strip()))
            continue
        total = _r2(_d(p.get("total")))
        sb.table("pedidos").update({
            "no_factura": no_factura.strip(),
            "facturado_con_iva": float(total),
            "facturado_sin_iva": float(_r2(importe_sin_iva(total)[0])),
        }).eq("numero_pedido", num).execute()
        guardados.append(p)
    return guardados, ya_tenian


# ----------------------------------------------------------------------
# Pestaña
# ----------------------------------------------------------------------
def _lector_constancia(elegidos):
    """Lee la Constancia de Situación Fiscal (PDF) en memoria y muestra los datos para copiar. No guarda nada."""
    ss = st.session_state
    clave = "_".join(str(p["numero_pedido"]) for p in elegidos) + f"_{ss.get('pf_ver', 0)}"
    with st.expander("📄 Leer Constancia de Situación Fiscal del cliente"):
        st.caption("Sube el PDF de la constancia que mandó el cliente: se lee en memoria para mostrarte los datos y "
                   "no se guarda en ningún lado (ni el archivo ni los datos).")
        archivo = st.file_uploader("Constancia (PDF)", type=["pdf"], key=f"pf_csf_{clave}",
                                   label_visibility="collapsed")
        if archivo is None:
            return
        try:
            from csf import leer_constancia
            d = leer_constancia(archivo.getvalue())
        except Exception as ex:
            st.error(f"No se pudo leer el PDF: {ex}")
            return
        pares = [("RFC", d.get("rfc")),
                 ("Nombre / razón social", d.get("nombre")),
                 ("Código postal", d.get("cp"))]
        for clave_r, nombre_r in d.get("regimenes", []):
            pares.append((f"Régimen fiscal {clave_r}", f"{clave_r} - {nombre_r}"))
        if len(elegidos) == 1 and elegidos[0].get("uso_cfdi"):
            pares.append(("Uso de CFDI (del pedido)", elegidos[0]["uso_cfdi"]))
        visibles = [x for x in pares if x[1]]
        if visibles:
            if d.get("tipo"):
                st.caption(f"Persona {d['tipo']}. En CFDI 4.0 el nombre va exactamente así, sin «S.A. de C.V.» ni "
                           "otro régimen de capital.")
            components.html(html_datos(visibles), height=46 + 34 * len(visibles))
        for aviso in d.get("avisos", []):
            st.warning(aviso)
        nombre_ped = (elegidos[0].get("cliente") or "").strip().upper() if len(elegidos) == 1 else ""
        if nombre_ped and d.get("nombre") and nombre_ped not in d["nombre"] and d["nombre"] not in nombre_ped:
            st.info(f"El nombre del pedido («{elegidos[0]['cliente']}») no coincide con el de la constancia: "
                    "revisa que sea la constancia correcta.")


def _consulta_facturados(sb, pedidos):
    """Vista de solo lectura de pedidos ya facturados, agrupados por N° de factura."""
    por_factura = {}
    for p in pedidos:
        por_factura.setdefault(p["no_factura"].strip(), []).append(p)
    items = items_de(sb, [p["numero_pedido"] for p in pedidos])
    for nf, peds in por_factura.items():
        nums = ", ".join(str(p["numero_pedido"]) for p in peds)
        total = sum((_r2(_d(p["total"])) for p in peds), Decimal("0"))
        st.markdown(f"🔒 **Factura {nf}** · pedido(s) {nums} · {peds[0]['cliente']}"
                    + (" y otros" if len({p['cliente'] for p in peds}) > 1 else "") + f" · ${total:,.2f}")
        with st.expander("Ver partidas (solo consulta)"):
            partidas = armar_partidas(peds, items, {p["numero_pedido"]: p.get("concepto") or DESGLOSAR for p in peds})
            if partidas:
                alto = 90 + 38 * len(partidas) + 16 * sum(len(x["avisos"]) for x in partidas)
                components.html(html_partidas(partidas), height=min(alto, 700), scrolling=alto > 700)
            st.caption("Partidas reconstruidas con el concepto guardado en el pedido. Si al facturar se cambió el "
                       "concepto, la factura real puede verse distinta. Para corregir el N° de factura usa "
                       "«Pedidos» → «Facturación y efectivo».")


def tab_por_facturar(sb, user):
    ss = st.session_state
    msg = ss.pop("pf_msg", None)
    if msg:
        st.success(msg)

    c1, c2 = st.columns([1, 3])
    desde = c1.date_input("Pedidos desde", value=dt.date.today() - dt.timedelta(days=90),
                          format="DD/MM/YYYY", key="pf_desde")
    ver = c2.radio("Mostrar", ["Todos", "Pendientes", "Facturados"], horizontal=True, key="pf_ver_filtro")
    try:
        todos = pedidos_con_factura(sb, desde)
    except Exception as ex:
        st.error(f"No se pudieron cargar los pedidos: {ex}")
        return

    pend_all = [p for p in todos if not facturado(p)]
    fact_all = [p for p in todos if facturado(p)]
    m1, m2, m3 = st.columns(3)
    m1.metric("POR FACTURAR", len(pend_all), help="Llevan datos de facturación y aún no tienen N° de factura.")
    m2.metric("MONTO POR FACTURAR", f"${sum(_d(p['total']) for p in pend_all):,.2f}")
    m3.metric("YA FACTURADOS", len(fact_all))

    lista = {"Todos": todos, "Pendientes": pend_all, "Facturados": fact_all}[ver]
    if not lista:
        st.info("No hay pedidos para mostrar con ese filtro." if todos else
                "No hay pedidos con datos de facturación en ese periodo.")
        return

    st.caption("Marca uno o varios pendientes para facturarlos (varios = factura mostrador). "
               "Los ya facturados solo se pueden consultar.")
    df = pd.DataFrame([{
        "SELECCIONAR": False,
        "ESTADO": "✅ Facturado" if facturado(p) else "🟡 Pendiente",
        "N° FACTURA": (p.get("no_factura") or "").strip() or "—",
        "N° PEDIDO": p["numero_pedido"], "FECHA": p["fecha"], "CLIENTE": p["cliente"],
        "CONCEPTO": p.get("concepto") or "—", "USO CFDI": p.get("uso_cfdi") or "—", "TOTAL": float(_d(p["total"])),
    } for p in lista])
    ed = st.data_editor(
        df, hide_index=True, use_container_width=True, key=f"pf_tabla_{ver}_{ss.get('pf_ver', 0)}",
        disabled=[c for c in df.columns if c != "SELECCIONAR"],
        column_config={"SELECCIONAR": st.column_config.CheckboxColumn("✔", width="small"),
                       "TOTAL": st.column_config.NumberColumn("TOTAL", format="$%.2f")})
    sel_num = set(ed.loc[ed["SELECCIONAR"], "N° PEDIDO"].tolist())
    if not sel_num:
        st.info("Marca en la tabla los pedidos que vas a facturar o consultar.")
        return

    # ---- Ya facturados: solo consulta ----
    ya = [p for p in lista if p["numero_pedido"] in sel_num and facturado(p)]
    elegidos = [p for p in lista if p["numero_pedido"] in sel_num and not facturado(p)]
    if ya:
        st.divider()
        if elegidos:
            st.warning("Marcaste pedidos ya facturados: no entran en la nueva factura, solo se muestran para consulta.")
        _consulta_facturados(sb, ya)
    if not elegidos:
        return

    st.divider()
    clientes = sorted({(p["cliente"] or "").strip() for p in elegidos})
    if len(elegidos) == 1:
        st.subheader(f"Factura del pedido {elegidos[0]['numero_pedido']}")
    else:
        st.subheader(f"Factura mostrador · {len(elegidos)} pedidos")
        if len(clientes) > 1:
            st.caption("Los pedidos son de clientes distintos: revisa a nombre de quién va la factura.")

    # Concepto por pedido (se puede cambiar aquí sin modificar el pedido)
    concepto_de = {}
    with st.expander("Concepto de facturación", expanded=len(elegidos) == 1 or
                     any(not p.get("concepto") for p in elegidos)):
        for p in elegidos:
            actual = p.get("concepto") if p.get("concepto") in CONCEPTOS else DESGLOSAR
            concepto_de[p["numero_pedido"]] = st.selectbox(
                f"Pedido {p['numero_pedido']} — {p['cliente']}", CONCEPTOS, index=CONCEPTOS.index(actual),
                key=f"pf_conc_{p['numero_pedido']}")
        con_modelo = st.checkbox("Incluir N° de modelo en la descripción (al desglosar)", key="pf_modelo")

    # Datos del cliente
    if len(elegidos) == 1:
        p = elegidos[0]
        datos = html_datos([
            ("Cliente / razón social", p.get("cliente")), ("Uso de CFDI", p.get("uso_cfdi")),
            ("E-mail (factura)", p.get("email_facturacion") or p.get("email")),
            ("Teléfono (factura)", p.get("telefono_facturacion") or p.get("telefono")),
            ("Notas de facturación", p.get("notas_facturacion"))])
        if datos:
            st.markdown("**Datos del cliente**")
            components.html(datos, height=46 + 34 * sum(1 for v in [
                p.get("cliente"), p.get("uso_cfdi"), p.get("email_facturacion") or p.get("email"),
                p.get("telefono_facturacion") or p.get("telefono"), p.get("notas_facturacion")] if v))

    # Constancia de situación fiscal (se lee en memoria, no se guarda)
    _lector_constancia(elegidos)

    # Partidas
    items = items_de(sb, [p["numero_pedido"] for p in elegidos])
    sin_items = [p["numero_pedido"] for p in elegidos if not items.get(p["numero_pedido"])]
    if sin_items:
        st.warning(f"Pedido(s) sin productos registrados: {', '.join(map(str, sin_items))}.")
    partidas = armar_partidas(elegidos, items, concepto_de, con_modelo)
    if not partidas:
        st.warning("No hay partidas con importe para facturar.")
        return
    st.markdown("**Partidas para CONTPAQi**")
    alto = 90 + 38 * len(partidas) + 16 * sum(len(x["avisos"]) for x in partidas)
    components.html(html_partidas(partidas), height=min(alto, 900), scrolling=alto > 900)
    nombre_xlsx = "factura_pedido_" + "_".join(str(p["numero_pedido"]) for p in elegidos)[:60] + ".xlsx"
    st.download_button("📥 Descargar Excel para CONTPAQi", excel_contpaqi(partidas), file_name=nombre_xlsx,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       type="primary", key=f"pf_xlsx_{nombre_xlsx}")
    st.caption("En CONTPAQi: elige el cliente → «Importar/Exportar Excel» → este archivo. "
               "Después elige **IVA 16%** en «Tipo Impuesto» de cada renglón y revisa que el total coincida.")

    total_ped = sum((_r2(_d(p["total"])) for p in elegidos), Decimal("0"))
    total_fac = sum((x["total"] for x in partidas), Decimal("0"))
    if total_ped != total_fac:
        st.warning(f"La suma de los pedidos (${total_ped:,.2f}) no coincide con la de las partidas "
                   f"(${total_fac:,.2f}). Revisa los productos del pedido.")
    st.caption("Los precios del pedido ya incluyen IVA; aquí se muestran sin IVA y ajustados para que el total "
               "timbrado coincida al centavo con lo cobrado.")

    # Guardar N° de factura
    st.divider()
    st.markdown("**Cuando ya esté timbrada en CONTPAQi**")
    no_fac = st.text_input("N° DE FACTURA (serie y folio)", key="pf_nofac", placeholder="Ej. A-1234")
    q_v, q_c = codigos.pedir_identidad(sb, "pf_ident")
    etiqueta = "Guardar N° de factura" + (f" en los {len(elegidos)} pedidos" if len(elegidos) > 1 else "")
    if st.button(etiqueta, type="primary", key="pf_guardar"):
        if not no_fac.strip():
            st.error("Escribe el N° de factura.")
        elif quien := codigos.confirmar(sb, q_v, q_c):
            try:
                guardados, ya_tenian = marcar_facturados(sb, elegidos, no_fac)
            except Exception as ex:
                st.error(f"No se pudo guardar: {ex}")
                return
            for p in guardados:
                codigos.registrar(sb, quien, "Registró factura", "pedido", p["numero_pedido"],
                                  f"factura {no_fac.strip()} · total {p['total']}")
            partes = []
            if guardados:
                partes.append(f"Factura {no_fac.strip()} guardada en pedido(s) "
                              f"{', '.join(str(p['numero_pedido']) for p in guardados)}.")
            if ya_tenian:
                partes.append("No se tocaron porque ya estaban facturados: "
                              + ", ".join(f"{n} (factura {f})" for n, f in ya_tenian) + ".")
            ss["pf_msg"] = " ".join(partes)
            ss["pf_ver"] = ss.get("pf_ver", 0) + 1
            ss.pop("pf_nofac", None)
            st.rerun()
